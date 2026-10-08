from datetime import datetime
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import Device
from app.reports import REPORT_HEADERS, build_equipment_pdf, build_equipment_xlsx


def report_rows():
    return [{
        "folder": "ЦОД / Стойка 1",
        "name": "Коммутатор ядра",
        "device_type": "Коммутатор",
        "ip_address": "192.0.2.10",
        "vendor": "МикроЭлектроника",
        "model": "МС-48 (рев. R2)",
        "installed_version": "1.0.0",
        "available_version": "1.2.0",
        "version_source": "SNMP",
        "status": "Есть обновление",
        "last_checked_at": datetime(2026, 10, 8, 12, 30),
        "last_checked_at_display": "08.10.2026 12:30",
        "description": "Основной коммутатор",
    }]


def test_excel_report_is_editable_filtered_and_preserves_cyrillic():
    content = build_equipment_xlsx(report_rows(), datetime(2026, 10, 8, 13, 0))
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Оборудование"]
    assert sheet.freeze_panes == "A7"
    assert [sheet.cell(6, column).value for column in range(1, 14)] == list(REPORT_HEADERS)
    assert sheet["C7"].value == "Коммутатор ядра"
    assert sheet["H7"].value == "1.0.0" and sheet["I7"].value == "1.2.0"
    assert sheet["K7"].value == "Есть обновление"
    assert sheet.tables["EquipmentReport"].ref == "A6:M7"


def test_report_content_cannot_inject_excel_formulas_or_pdf_markup():
    rows = report_rows()
    rows[0]["name"] = "=HYPERLINK(\"https://example.test\")"
    rows[0]["description"] = "<b>не HTML</b> & обычный текст"
    workbook = load_workbook(BytesIO(build_equipment_xlsx(rows, datetime(2026, 10, 8, 13, 0))))
    assert workbook["Оборудование"]["C7"].data_type == "s"
    assert workbook["Оборудование"]["C7"].value.startswith("'")
    assert build_equipment_pdf(rows, datetime(2026, 10, 8, 13, 0)).startswith(b"%PDF")


def test_pdf_report_is_readable_and_preserves_cyrillic():
    content = build_equipment_pdf(report_rows(), datetime(2026, 10, 8, 13, 0))
    assert content.startswith(b"%PDF")
    assert len(content) > 20_000


def test_report_download_endpoints_and_dashboard_actions():
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        device = Device(name="Report endpoint device", ip_address="192.0.2.44", vendor="Vendor", model="Model", installed_version="2.0", status="Актуально")
        db.add(device);db.commit();device_id=device.id
    try:
        with TestClient(app) as client:
            xlsx = client.get("/reports/equipment.xlsx")
            pdf = client.get("/reports/equipment.pdf")
            dashboard = client.get("/")
        assert xlsx.status_code == 200 and xlsx.content.startswith(b"PK")
        assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert "attachment" in xlsx.headers["content-disposition"] and ".xlsx" in xlsx.headers["content-disposition"]
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert pdf.headers["content-type"].startswith("application/pdf")
        assert "/reports/equipment.pdf" in dashboard.text and "/reports/equipment.xlsx" in dashboard.text
    finally:
        with SessionLocal() as db:
            device=db.get(Device,device_id)
            if device:db.delete(device);db.commit()


def test_dashboard_report_menu_matches_project_styles():
    root = Path(__file__).parents[1]
    html = (root / "app/templates/dashboard.html").read_text(encoding="utf-8")
    css = (root / "app/static/report-export.css").read_text(encoding="utf-8")
    assert "Выгрузить отчёт" in html and "Можно редактировать" in html
    assert ".report-export>div" in css and "#121d2b" in css
