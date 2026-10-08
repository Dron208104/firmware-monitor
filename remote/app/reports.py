"""Printable and editable equipment reports."""

from __future__ import annotations

from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any


REPORT_HEADERS = (
    "№", "Каталог", "Устройство", "Тип", "IP-адрес", "Производитель", "Модель",
    "Установленная версия", "Доступная версия", "Источник версии", "Статус",
    "Последняя проверка", "Комментарий",
)


def _summary(rows: list[dict[str, Any]]) -> dict[str, int]:
    updates = {"Доступно обновление", "Есть обновление"}
    current = {"Актуальная версия", "Актуально"}
    return {
        "total": len(rows),
        "updates": sum(row["status"] in updates for row in rows),
        "current": sum(row["status"] in current for row in rows),
        "review": sum(row["status"] not in updates | current for row in rows),
    }


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def build_equipment_xlsx(rows: list[dict[str, Any]], generated_at) -> bytes:
    """Build an editable Excel register with filters and frozen headers."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Оборудование"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws["A2"] = "Отчёт по оборудованию и версиям прошивок"
    ws["A2"].font = Font(name="Arial", size=15, bold=True, color="183153")
    ws["A3"] = f"Сформирован: {generated_at.strftime('%d.%m.%Y %H:%M')}"
    ws["A3"].font = Font(name="Arial", size=10, italic=True, color="64748B")

    summary = _summary(rows)
    summary_values = (
        ("Всего устройств", summary["total"]), ("Есть обновления", summary["updates"]),
        ("Актуальны", summary["current"]), ("Требуют проверки", summary["review"]),
    )
    for column, (label, value) in zip((1, 4, 7, 10), summary_values):
        ws.cell(4, column, label).font = Font(name="Arial", size=10, bold=True, color="334155")
        value_cell = ws.cell(4, column + 1, value)
        value_cell.font = Font(name="Arial", size=12, bold=True, color="2563EB")

    header_row = 6
    for column, heading in enumerate(REPORT_HEADERS, start=1):
        cell = ws.cell(header_row, column, heading)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183153")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    status_fills = {
        "Актуально": ("E8F7EF", "167A52"), "Актуальная версия": ("E8F7EF", "167A52"),
        "Есть обновление": ("FFF4DD", "A66400"), "Доступно обновление": ("FFF4DD", "A66400"),
    }
    thin = Side(style="thin", color="DCE5EF")
    for index, row in enumerate(rows, start=1):
        excel_row = header_row + index
        values = (
            index, row["folder"], row["name"], row["device_type"], row["ip_address"],
            row["vendor"], row["model"], row["installed_version"], row["available_version"],
            row["version_source"], row["status"], row["last_checked_at"], row["description"],
        )
        for column, value in enumerate(values, start=1):
            cell = ws.cell(excel_row, column, _excel_safe(value))
            cell.font = Font(name="Arial", size=10, color="172033")
            cell.alignment = Alignment(vertical="top", wrap_text=column in {2, 3, 7, 11, 13})
            cell.border = Border(bottom=thin)
        ws.cell(excel_row, 12).number_format = "dd.mm.yyyy hh:mm"
        status_cell = ws.cell(excel_row, 11)
        fill, color = status_fills.get(row["status"], ("EEF3FA", "48627F"))
        status_cell.fill = PatternFill("solid", fgColor=fill)
        status_cell.font = Font(name="Arial", size=10, bold=True, color=color)
        ws.row_dimensions[excel_row].height = 32

    last_row = header_row + max(len(rows), 1)
    if rows:
        table = Table(displayName="EquipmentReport", ref=f"A{header_row}:M{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A{header_row}:M{header_row}"

    for column, width in enumerate((6, 20, 22, 18, 16, 18, 24, 20, 18, 17, 24, 20, 34), start=1):
        ws.column_dimensions[ws.cell(1, column).column_letter].width = width
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Firmware Monitor"
    ws.oddFooter.right.text = "Страница &P из &N"

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _find_pdf_fonts() -> tuple[str, str]:
    candidates = (
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
    )
    for regular, bold in candidates:
        if regular.is_file() and bold.is_file():
            return str(regular), str(bold)
    raise RuntimeError("Не найден шрифт с поддержкой кириллицы для PDF")


def build_equipment_pdf(rows: list[dict[str, Any]], generated_at) -> bytes:
    """Build a compact landscape PDF suitable for printing and archiving."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold = _find_pdf_fonts()
    if "FirmwareMonitor" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("FirmwareMonitor", regular))
        pdfmetrics.registerFont(TTFont("FirmwareMonitor-Bold", bold))

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=9*mm, rightMargin=9*mm, topMargin=9*mm, bottomMargin=11*mm, title="Отчёт по оборудованию и версиям прошивок", author="Firmware Monitor")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName="FirmwareMonitor-Bold", fontSize=16, leading=19, textColor=colors.HexColor("#183153"), alignment=TA_LEFT, spaceAfter=3*mm)
    meta = ParagraphStyle("ReportMeta", parent=styles["Normal"], fontName="FirmwareMonitor", fontSize=8, leading=11, textColor=colors.HexColor("#64748B"), alignment=TA_LEFT)
    summary_style = ParagraphStyle("ReportSummary", parent=meta, fontName="FirmwareMonitor-Bold", textColor=colors.HexColor("#334155"))
    header = ParagraphStyle("ReportHeader", parent=meta, fontName="FirmwareMonitor-Bold", fontSize=6.2, leading=7.4, textColor=colors.white)
    body = ParagraphStyle("ReportBody", parent=meta, fontSize=6.2, leading=7.6, textColor=colors.HexColor("#172033"))
    summary = _summary(rows)
    story = [
        Paragraph("Отчёт по оборудованию и версиям прошивок", title),
        Paragraph(f"Сформирован: {generated_at.strftime('%d.%m.%Y %H:%M')}", meta), Spacer(1, 2*mm),
        Paragraph(f"Всего устройств: {summary['total']} &nbsp;&nbsp;&nbsp; Есть обновления: {summary['updates']} &nbsp;&nbsp;&nbsp; Актуальны: {summary['current']} &nbsp;&nbsp;&nbsp; Требуют проверки: {summary['review']}", summary_style),
        Spacer(1, 3*mm),
    ]
    pdf_headers = REPORT_HEADERS[:-1]
    data = [[Paragraph(value, header) for value in pdf_headers]]
    for index, row in enumerate(rows, start=1):
        values = (str(index), row["folder"], row["name"], row["device_type"], row["ip_address"], row["vendor"], row["model"], row["installed_version"], row["available_version"], row["version_source"], row["status"], row["last_checked_at_display"])
        data.append([Paragraph(escape(str(value or "—")), body) for value in values])
    if not rows:
        data.append([Paragraph("Нет доступного оборудования", body)] + [""] * (len(pdf_headers)-1))
    table = Table(data, repeatRows=1, colWidths=[7*mm,22*mm,28*mm,22*mm,21*mm,23*mm,31*mm,24*mm,23*mm,20*mm,29*mm,23*mm])
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#183153")), ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),3), ("RIGHTPADDING",(0,0),(-1,-1),3),
        ("TOPPADDING",(0,0),(-1,0),5), ("BOTTOMPADDING",(0,0),(-1,0),5),
        ("TOPPADDING",(0,1),(-1,-1),4), ("BOTTOMPADDING",(0,1),(-1,-1),4),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F6F9FC")]),
        ("LINEBELOW",(0,0),(-1,-1),.35,colors.HexColor("#DCE5EF")),
    ]))
    story.append(table)

    def footer(canvas, document):
        canvas.saveState();canvas.setFont("FirmwareMonitor",7);canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(9*mm,6*mm,"Firmware Monitor")
        canvas.drawRightString(landscape(A4)[0]-9*mm,6*mm,f"Страница {document.page}")
        canvas.restoreState()

    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()
